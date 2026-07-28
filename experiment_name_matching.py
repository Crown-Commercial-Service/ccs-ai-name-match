"""Run the name-matching Excel truthset against a temporary local API server."""

import argparse
import ast
import json
import subprocess
import sys
import time
import urllib.parse
import urllib.request
from contextlib import contextmanager
from pathlib import Path

DEFAULT_TRUTHSET = Path("name-match-api truthset.xlsx")
DEFAULT_HOST = "127.0.0.1"
DEFAULT_PORT = 8000

# The supplied spreadsheet uses: input, options, true_match.
COLUMN_NAMES = {
    "input": ("input", "input_string", "source_name", "unmatched_name", "name_to_match", "query"),
    "expected match": (
        "true_match",
        "expected_match",
        "expected",
        "target_name",
        "matched_name",
        "correct_match",
        "ground_truth",
        "canonical_name",
    ),
    "candidates": ("options", "candidates", "candidate_list", "list_of_strings"),
}


def _http_get(url, timeout_s=60.0):
    """Return an HTTP GET response as (status code, text)."""
    try:
        with urllib.request.urlopen(url, timeout=timeout_s) as response:
            return response.status, response.read().decode("utf-8")
    except Exception as exc:
        raise RuntimeError(f"Request failed for {url}: {exc}") from exc


def match_string_via_api(
    input_string,
    list_of_strings,
    prompt_path=None,
    api_url=None,
    timeout_s=60.0,
    extra_query_params=None,
):
    """Call the local /match endpoint and return its selected candidate."""
    api_url = api_url or f"http://{DEFAULT_HOST}:{DEFAULT_PORT}/match"
    candidates = [candidate for candidate in list_of_strings if candidate != input_string]
    query = {"input_string": input_string, "candidates": candidates}

    if prompt_path:
        query["prompt_path"] = prompt_path
    if extra_query_params:
        query.update(extra_query_params)

    separator = "&" if "?" in api_url else "?"
    url = api_url + separator + urllib.parse.urlencode(query, doseq=True)
    status, response_text = _http_get(url, timeout_s)
    if not 200 <= status < 300:
        raise RuntimeError(f"Match API returned status {status}: {response_text}")

    try:
        response = json.loads(response_text)
        result = response.get("match") if isinstance(response, dict) else response
    except json.JSONDecodeError:
        result = response_text

    result = "" if result is None else str(result).strip()
    if not result or result.lower() in {"none", "null", "n/a", "na"}:
        return "None"
    return result if result in list_of_strings else "None"


def _normalise_header(value):
    return str(value or "").strip().lower().replace(" ", "_").replace("-", "_")


def _find_column(headers, label, required=True):
    """Find a column using the accepted names in COLUMN_NAMES."""
    for accepted_name in COLUMN_NAMES[label]:
        if accepted_name in headers:
            return headers.index(accepted_name)

    if required:
        raise ValueError(
            f"Could not find the {label} column. Found headers: {', '.join(headers)}. "
            f"Accepted names: {', '.join(COLUMN_NAMES[label])}"
        )
    return None


def _parse_candidates(value):
    """Turn an Excel options cell into a list of candidate names."""
    if value is None:
        return []
    if isinstance(value, (list, tuple)):
        return [str(item).strip() for item in value if str(item).strip()]

    text = str(value).strip()
    if not text:
        return []

    # Excel cells may contain a JSON list or a Python-style list.
    for parser in (json.loads, ast.literal_eval):
        try:
            parsed = parser(text)
            if isinstance(parsed, (list, tuple)):
                return [str(item).strip() for item in parsed if str(item).strip()]
        except (ValueError, SyntaxError, json.JSONDecodeError):
            continue

    separator = "\n" if "\n" in text else "|" if "|" in text else ";"
    return [item.strip() for item in text.split(separator) if item.strip()]


def load_truthset(path):
    """Read (input, true match, options) rows from the first Excel worksheet."""
    from openpyxl import load_workbook

    worksheet = load_workbook(path, read_only=True, data_only=True).active
    rows = worksheet.iter_rows(values_only=True)
    header_row = next(rows, None)
    if header_row is None:
        raise ValueError(f"Truthset is empty: {path}")

    headers = [_normalise_header(header) for header in header_row]
    input_column = _find_column(headers, "input")
    expected_column = _find_column(headers, "expected match")
    candidates_column = _find_column(headers, "candidates", required=False)

    truthset = []
    for row in rows:
        input_value = row[input_column] if input_column < len(row) else None
        expected_value = row[expected_column] if expected_column < len(row) else None
        if input_value is None or expected_value is None:
            continue

        candidates = None
        if candidates_column is not None and candidates_column < len(row):
            candidates = _parse_candidates(row[candidates_column])

        truthset.append((str(input_value).strip(), str(expected_value).strip(), candidates))

    if not truthset:
        raise ValueError(f"Truthset contains no usable data rows: {path}")
    return truthset


def _wait_for_server(base_url, process, timeout_s=30):
    """Wait until Uvicorn's health endpoint is ready."""
    deadline = time.monotonic() + timeout_s
    while time.monotonic() < deadline:
        if process.poll() is not None:
            raise RuntimeError(f"Uvicorn exited with code {process.returncode}")
        try:
            status, _ = _http_get(f"{base_url}/health", timeout_s=1)
            if status == 200:
                return
        except RuntimeError:
            time.sleep(0.2)
    raise TimeoutError("Uvicorn did not start within 30 seconds")


@contextmanager
def local_uvicorn(host=DEFAULT_HOST, port=DEFAULT_PORT):
    """Start app.main:app locally, then stop it when the experiment finishes."""
    command = [
        sys.executable,
        "-m",
        "uvicorn",
        "app.main:app",
        "--host",
        host,
        "--port",
        str(port),
    ]
    process = subprocess.Popen(command, cwd=Path(__file__).resolve().parent)
    base_url = f"http://{host}:{port}"

    try:
        _wait_for_server(base_url, process)
        yield f"{base_url}/match"
    finally:
        if process.poll() is None:
            process.terminate()
            try:
                process.wait(timeout=5)
            except subprocess.TimeoutExpired:
                process.kill()
                process.wait()


def run_experiment(truthset_path, api_url, prompt_path=None):
    """Match every truthset row, print each result, and print overall accuracy."""
    rows = load_truthset(truthset_path)
    all_expected_matches = list(dict.fromkeys(expected for _, expected, _ in rows))
    correct = 0

    for excel_row, (input_name, expected, row_candidates) in enumerate(rows, start=2):
        candidates = row_candidates or all_expected_matches
        if expected not in candidates:
            candidates = candidates + [expected]

        actual = match_string_via_api(
            input_name,
            candidates,
            prompt_path=prompt_path,
            api_url=api_url,
        )
        passed = actual == expected
        correct += passed
        print(
            f"Row {excel_row-1}: {'PASS' if passed else 'FAIL'} | "
            f"input={input_name!r} | expected={expected!r} | actual={actual!r}"
        )

    accuracy = correct / len(rows)
    print(f"Accuracy: {accuracy:.2%} ({correct}/{len(rows)})")
    return accuracy


def main():
    parser = argparse.ArgumentParser(description="Run the name-matching truthset locally")
    parser.add_argument("--truthset", type=Path, default=DEFAULT_TRUTHSET)
    parser.add_argument("--host", default=DEFAULT_HOST)
    parser.add_argument("--port", type=int, default=DEFAULT_PORT)
    parser.add_argument("--prompt-path")
    args = parser.parse_args()

    if not args.truthset.is_file():
        raise FileNotFoundError(f"Truthset not found: {args.truthset}")

    with local_uvicorn(args.host, args.port) as api_url:
        run_experiment(args.truthset, api_url, args.prompt_path)


if __name__ == "__main__":
    main()
